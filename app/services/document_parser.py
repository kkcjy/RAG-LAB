# 文档解析：PDF / Word / Markdown / TXT → 纯文本
#   PDF：PyMuPDF 优先，回退 pdfplumber；两者都无文字时自动转 OCR（扫描件）
#   OCR：本地 RapidOCR（PP-OCRv4），模型内置，CPU 推理，无需联网
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path

logger = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    text: str
    page_count: int = 0
    meta: dict | None = None


class UnsupportedFileType(Exception):
    pass


# ---------------------------------------------------------------------------
# 各类型解析器
# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------
# OCR（扫描件 PDF / 图片型文档）
# ---------------------------------------------------------------------------
_ocr_engine = None


def _get_ocr_engine():
    """懒加载 OCR 引擎（首次调用时初始化，约 1 秒）。

    rapidocr_onnxruntime 模型内置于包中，无需联网下载。
    """
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR

        logger.info("正在初始化 OCR 引擎（RapidOCR / PP-OCRv4）")
        _ocr_engine = RapidOCR()
    return _ocr_engine


def _ocr_image_bytes(img_bytes: bytes) -> str:
    """对单张图片做 OCR，返回识别文本（按行拼接）"""
    import io

    import numpy as np
    from PIL import Image

    engine = _get_ocr_engine()
    img = Image.open(io.BytesIO(img_bytes)).convert("RGB")
    result, _ = engine(np.array(img))
    if not result:
        return ""
    # result 形如 [[box, text, score], ...]
    return "\n".join(str(line[1]) for line in result if len(line) > 1)


def _ocr_pdf(path: Path) -> tuple[str, int, bool]:
    """对扫描件 PDF 做 OCR。

    返回 (文本, 已处理页数, 是否因超出上限被截断)。
    用 PyMuPDF 渲染页面为位图（无需额外系统依赖，比 pdftoppm 少一次进程调用）。
    """
    import fitz  # PyMuPDF

    from app.core.config import settings

    max_pages = settings.ocr_max_pages
    dpi = settings.ocr_dpi
    parts: list[str] = []

    with fitz.open(str(path)) as doc:
        total = doc.page_count
        limit = min(total, max_pages)
        for i in range(limit):
            page = doc[i]
            # 渲染为 PNG 位图
            pix = page.get_pixmap(dpi=dpi)
            text = _ocr_image_bytes(pix.tobytes("png"))
            if text.strip():
                parts.append(text)
            logger.debug("OCR 第 %d/%d 页完成，识别 %d 字符", i + 1, limit, len(text))

    return "\n\n".join(parts), limit, total > max_pages


def _parse_pdf(path: Path) -> ParsedDocument:
    """优先 PyMuPDF（快、稳），失败回退 pdfplumber；无文字则尝试 OCR。"""
    text_parts: list[str] = []
    page_count = 0

    try:
        import fitz  # PyMuPDF

        with fitz.open(str(path)) as doc:
            page_count = doc.page_count
            for page in doc:
                text_parts.append(page.get_text("text"))
        text = "\n".join(text_parts)
        if text.strip():
            return ParsedDocument(text=text, page_count=page_count, meta={"parser": "pymupdf"})
        logger.warning("PyMuPDF 未提取到文本，回退 pdfplumber: %s", path.name)
    except Exception as e:  # noqa: BLE001
        logger.warning("PyMuPDF 解析失败(%s)，回退 pdfplumber", e)

    # 回退
    try:
        import pdfplumber

        parts = []
        with pdfplumber.open(str(path)) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
        text = "\n".join(parts)
    except Exception as e:  # noqa: BLE001
        logger.warning("pdfplumber 解析失败(%s)", e)
        text = ""

    if text.strip():
        return ParsedDocument(text=text, page_count=page_count, meta={"parser": "pdfplumber"})

    # 两个解析器都读不到文字 —— 很可能是扫描件，尝试 OCR
    from app.core.config import settings

    if not settings.ocr_enabled:
        logger.warning("PDF 无文本且 OCR 已禁用: %s", path.name)
        return ParsedDocument(text="", page_count=page_count, meta={"parser": "none"})

    try:
        logger.info("PDF 无文本，尝试 OCR: %s", path.name)
        ocr_text, done_pages, truncated = _ocr_pdf(path)
        meta = {
            "parser": "rapidocr",
            "ocr_pages": done_pages,
            "ocr_truncated": truncated,
        }
        if truncated:
            logger.warning(
                "文档页数超过 OCR 上限，仅处理前 %d 页: %s", done_pages, path.name
            )
        return ParsedDocument(text=ocr_text, page_count=page_count or done_pages, meta=meta)
    except Exception as e:  # noqa: BLE001
        logger.error("OCR 失败(%s): %s", path.name, e)
        return ParsedDocument(text="", page_count=page_count, meta={"parser": "ocr_failed"})


def _parse_docx(path: Path) -> ParsedDocument:
    """python-docx：段落 + 表格 + 页眉页脚"""
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts: list[str] = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    # 表格按 Markdown 风格保留结构，便于模型理解
    for ti, table in enumerate(doc.tables):
        parts.append(f"\n[表格 {ti + 1}]")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            parts.append(" | ".join(cells))

    # 页眉页脚（部分文档正文在这）
    for section in doc.sections:
        for holder in (section.header, section.footer):
            for para in holder.paragraphs:
                if para.text.strip():
                    parts.append(para.text)

    return ParsedDocument(
        text="\n".join(parts),
        page_count=0,
        meta={"parser": "python-docx", "tables": len(doc.tables)},
    )


def _parse_markdown(path: Path) -> ParsedDocument:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return ParsedDocument(text=text, meta={"parser": "plain"})


def _parse_txt(path: Path) -> ParsedDocument:
    for enc in ("utf-8", "gbk", "gb18030", "latin-1"):
        try:
            text = path.read_text(encoding=enc)
            return ParsedDocument(text=text, meta={"parser": "plain", "encoding": enc})
        except (UnicodeDecodeError, LookupError):
            continue
    raise ValueError(f"无法解码文本文件: {path.name}")


def _parse_doc(path: Path) -> ParsedDocument:
    """老式 .doc —— 尝试 antiword / libreoffice，都没有则报错提示转 docx"""
    import shutil
    import subprocess

    if shutil.which("antiword"):
        out = subprocess.run(
            ["antiword", str(path)], capture_output=True, text=True, timeout=60
        )
        if out.returncode == 0 and out.stdout.strip():
            return ParsedDocument(text=out.stdout, meta={"parser": "antiword"})

    if shutil.which("libreoffice"):
        import tempfile

        with tempfile.TemporaryDirectory() as td:
            subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "docx", "--outdir", td, str(path)],
                capture_output=True,
                timeout=180,
            )
            converted = list(Path(td).glob("*.docx"))
            if converted:
                return _parse_docx(converted[0])

    raise UnsupportedFileType(
        ".doc 格式需要 antiword 或 libreoffice 支持，建议转为 .docx 后上传"
    )


_PARSERS = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".doc": _parse_doc,
    ".md": _parse_markdown,
    ".markdown": _parse_markdown,
    ".txt": _parse_txt,
}


def parse_document(path: str | Path) -> ParsedDocument:
    p = Path(path)
    ext = p.suffix.lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise UnsupportedFileType(f"不支持的文件类型: {ext}")

    logger.info("正在解析文档 %s (%s)", p.name, ext)
    result = parser(p)
    result.text = normalize_text(result.text)
    return result


# ---------------------------------------------------------------------------
def normalize_text(text: str) -> str:
    """清洗：去多余空白、修复断行、统一换行符"""
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # 全角空格 → 半角
    text = text.replace("\u3000", " ")
    # PDF/OCR 提取文本常夹杂 NUL 字节（\x00），PostgreSQL 的 TEXT 列按 UTF8
    # 编码存储会直接拒绝它，导致写分块时整条索引失败。顺带清掉其余无意义的
    # 控制字符（保留 \n 与 \t）。
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    # 压缩 3 个以上连续换行为 2 个（保留段落感）
    text = re.sub(r"\n{3,}", "\n\n", text)
    # 行内多余空格
    text = re.sub(r"[ \t]{2,}", " ", text)
    # 去掉行首行尾空白
    lines = [ln.strip() for ln in text.split("\n")]
    return "\n".join(lines).strip()


# ---------------------------------------------------------------------------
def split_text(
    text: str,
    chunk_size: int = 800,
    chunk_overlap: int = 150,
) -> list[str]:
    """
    中文友好的递归分块。
    优先在段落/句号/换行处切分，避免把句子切断。
    """
    if not text.strip():
        return []

    # 参数防护：chunk_size / chunk_overlap 可由管理员在运行时通过
    # PUT /api/admin/config 修改，且没有校验。overlap >= chunk_size 时硬切步长为
    # 0（range 抛 ValueError）或负数（range 为空，整段内容被静默丢弃）。
    # 上限取 chunk_size // 2：更大的 overlap 虽不再报错，但步长趋近 1，
    # 会把一段文本切成大量近乎重复的碎块、撑爆索引。
    if chunk_size <= 0:
        chunk_size = 800
    chunk_overlap = max(0, min(chunk_overlap, chunk_size // 2))

    separators = ["\n\n", "\n", "。", "！", "？", "；", ".", "!", "?", ";", "，", ",", " ", ""]

    chunks: list[str] = []

    def _split(segment: str, seps: list[str]) -> None:
        if len(segment) <= chunk_size:
            if segment.strip():
                chunks.append(segment.strip())
            return

        if not seps:
            # 无分隔符可用，硬切
            for i in range(0, len(segment), chunk_size - chunk_overlap):
                piece = segment[i:i + chunk_size]
                if piece.strip():
                    chunks.append(piece.strip())
            return

        sep = seps[0]
        rest = seps[1:]

        if sep == "":
            for i in range(0, len(segment), chunk_size - chunk_overlap):
                piece = segment[i:i + chunk_size]
                if piece.strip():
                    chunks.append(piece.strip())
            return

        parts = segment.split(sep)
        buffer = ""
        for part in parts:
            candidate = buffer + sep + part if buffer else part
            if len(candidate) <= chunk_size:
                buffer = candidate
            else:
                if buffer:
                    _split(buffer, rest)
                buffer = part
        if buffer:
            _split(buffer, rest)

    _split(text, separators)

    # 合并过短的块（避免碎片）
    merged: list[str] = []
    min_len = max(50, chunk_size // 10)
    for c in chunks:
        if merged and len(c) < min_len and len(merged[-1]) + len(c) <= chunk_size:
            merged[-1] = merged[-1] + "\n" + c
        else:
            merged.append(c)

    # 应用 chunk_overlap：把前一块的尾部并入当前块开头，避免跨块语义被切断。
    # 走自然分隔符路径时上面的递归切分并不会用到 overlap，只有无分隔符的硬切
    # 回退才会用，所以此前配置的 chunk_overlap 对普通文档完全不生效。
    # 拼接后仍不超过 chunk_size：空间不足时按剩余空间截取重叠长度。
    if chunk_overlap > 0 and len(merged) > 1:
        overlapped = [merged[0]]
        for prev, cur in zip(merged, merged[1:]):
            # -1 给拼接用的换行留位，保证整块仍不超过 chunk_size
            room = chunk_size - len(cur) - 1
            if room > 0:
                tail = prev[-min(chunk_overlap, room):].strip()
                if tail:
                    cur = f"{tail}\n{cur}"
            overlapped.append(cur)
        merged = overlapped

    return merged
